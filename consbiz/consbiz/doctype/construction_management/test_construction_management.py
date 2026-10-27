# Copyright (c) 2026, Siricons and Contributors
# See license.txt

import frappe
from frappe.tests import IntegrationTestCase
from frappe.utils import getdate

from consbiz.consbiz.doctype.construction_management.construction_management import (
	add_working_days,
	count_working_days,
	plan_progress,
	target_progress,
)

EXTRA_TEST_RECORD_DEPENDENCIES = []
IGNORE_TEST_RECORD_DEPENDENCIES = ["BOQ", "Customer", "Company", "Holiday List", "UOM", "Project"]


class IntegrationTestConstructionManagement(IntegrationTestCase):
	def make_schedule(self, tasks):
		doc = frappe.new_doc("Construction Management")
		doc.start_date = "2026-11-02"
		for task in tasks:
			doc.append("tasks", task)
		doc.flags.ignore_links = True
		# ไม่ insert: ทดสอบเฉพาะการคำนวณใน validate (ข้ามเช็ก BOQ ที่ใช้งานซ้ำ)
		doc.status = "Inactive"
		doc.run_method("validate")
		return doc

	def test_structure_weight_and_rollup(self):
		doc = self.make_schedule(
			[
				{"wbs": "1", "item_name": "งานเตรียมการ"},
				{"wbs": "1.1", "amount": 1000, "start_date": "2026-11-02", "end_date": "2026-11-05"},
				{"wbs": "1.2", "amount": 3000, "start_date": "2026-11-04", "duration": 3},
				{"wbs": "2", "item_name": "งานโครงสร้าง"},
				{"wbs": "2.1", "amount": 6000, "start_date": "2026-11-10", "end_date": "2026-11-10"},
			]
		)
		g1, t11, t12, g2, _t21 = doc.tasks

		self.assertEqual([r.indent for r in doc.tasks], [0, 1, 1, 0, 1])
		self.assertEqual([r.is_group for r in doc.tasks], [1, 0, 0, 1, 0])
		self.assertEqual(doc.total_amount, 10000)
		self.assertAlmostEqual(t11.weight, 10)
		self.assertAlmostEqual(g1.weight, 40)
		self.assertAlmostEqual(g2.weight, 60)
		self.assertEqual(g1.amount, 4000)

		self.assertEqual(t11.duration, 4)
		self.assertEqual(getdate(t12.end_date), getdate("2026-11-06"))
		self.assertEqual(getdate(g1.start_date), getdate("2026-11-02"))
		self.assertEqual(getdate(g1.end_date), getdate("2026-11-06"))
		self.assertEqual(g1.duration, 5)

	def test_contract_duration(self):
		doc = self.make_schedule(
			[{"wbs": "1", "amount": 1, "start_date": "2026-11-02", "end_date": "2026-11-05"}]
		)
		self.assertEqual(doc.duration, 0)
		doc.finish_date = "2026-11-30"
		doc.run_method("validate")
		# 2 - 30 พ.ย. นับรวมวันแรก = 29 วัน
		self.assertEqual(doc.duration, 29)
		doc.finish_date = "2026-11-01"
		with self.assertRaises(frappe.ValidationError):
			doc.run_method("validate")

	def test_end_before_start(self):
		with self.assertRaises(frappe.ValidationError):
			self.make_schedule(
				[{"wbs": "1", "amount": 1, "start_date": "2026-11-05", "end_date": "2026-11-04"}]
			)

	def test_plan_progress(self):
		# น้ำหนัก 50/50: งานแรกกระจายเท่ากัน 10 วัน งานที่สองกำหนดเองรายสัปดาห์ในเดือน 10/20/30/40
		doc = self.make_schedule(
			[
				{"wbs": "1", "amount": 1000, "start_date": "2026-11-02", "end_date": "2026-11-11"},
				{
					"wbs": "2",
					"amount": 1000,
					"start_date": "2026-11-01",
					"end_date": "2026-11-30",
					"plan_distribution": {"unit": "mweek", "weights": [10, 20, 30, 40]},
				},
			]
		)
		self.assertEqual(plan_progress(doc, "2026-10-31"), 0)
		# 50 x 5/10 + 50 x (10% x 6/7)
		self.assertAlmostEqual(plan_progress(doc, "2026-11-06"), 25 + 50 * 0.1 * 6 / 7)
		self.assertAlmostEqual(plan_progress(doc, "2026-11-14"), 50 + 50 * 0.3)
		# ช่วงสุดท้าย 22-30 พ.ย. (9 วัน) ทำไป 4 วัน
		self.assertAlmostEqual(plan_progress(doc, "2026-11-25"), 50 + 50 * (0.6 + 0.4 * 4 / 9))
		self.assertAlmostEqual(plan_progress(doc, "2026-12-01"), 100)

		# นับวันทำงาน: วันหยุดไม่มี % แผน
		doc.tasks[1].plan_distribution = None
		doc.use_working_days = 1
		doc.get_holidays = lambda: {getdate("2026-11-03")}
		self.assertAlmostEqual(plan_progress(doc, "2026-11-03"), 50 / 9 + 50 * 2 / 29)

	def test_target_progress(self):
		doc = self.make_schedule(
			[{"wbs": "1", "amount": 1000, "start_date": "2026-11-02", "end_date": "2026-11-11"}]
		)
		# ยังไม่ตั้ง Baseline: เทียบกับแผนปัจจุบัน
		value, basis = target_progress(doc, "2026-11-06")
		self.assertAlmostEqual(value, 50)
		self.assertEqual(basis, "Current Plan")

		# Baseline 2-6 พ.ย. ส่วนแผนปัจจุบันขยายถึง 11 พ.ย.
		row = doc.tasks[0]
		row.baseline_start_date, row.baseline_end_date, row.baseline_weight = "2026-11-02", "2026-11-06", 100
		value, basis = target_progress(doc, "2026-11-04")
		self.assertAlmostEqual(value, 60)
		self.assertEqual(basis, "Baseline")
		self.assertAlmostEqual(plan_progress(doc, "2026-11-04"), 30)

	def make_linked(self, tasks, holidays=None):
		"""แผนที่แถวมีชื่อ (ใช้อ้างเป็นงานก่อนหน้า) predecessors ใส่เป็น list และ holidays แทนปฏิทินวันหยุด
		คืน {ชื่อแถว: แถว} เอกสารอยู่ที่ self.linked_doc"""
		doc = frappe.new_doc("Construction Management")
		doc.start_date = "2026-11-02"
		doc.status = "Inactive"
		doc.flags.ignore_links = True
		for task in tasks:
			doc.append("tasks", {**task, "predecessors": frappe.as_json(task.get("predecessors") or [])})
		if holidays:
			doc.use_working_days = 1
			doc.get_holidays = lambda: {getdate(d) for d in holidays}
		doc.run_method("validate")
		self.linked_doc = doc
		return {row.name: row for row in doc.tasks}

	def assert_dates(self, row, start, end):
		self.assertEqual((getdate(row.start_date), getdate(row.end_date)), (getdate(start), getdate(end)))

	def test_dependency_types(self):
		# งานก่อนหน้า 2-4 พ.ย. งานที่ตามมายาว 2 วัน ตั้งไว้ก่อนวันที่ความสัมพันธ์ยอม จึงถูกดันออก
		expected = {
			"FS": ("2026-11-05", "2026-11-06"),
			"SS": ("2026-11-02", "2026-11-03"),
			"FF": ("2026-11-03", "2026-11-04"),
			"SF": ("2026-10-31", "2026-11-01"),
		}
		for link_type, (start, end) in expected.items():
			rows = self.make_linked(
				[
					{
						"name": "a",
						"wbs": "1",
						"amount": 1,
						"start_date": "2026-11-02",
						"end_date": "2026-11-04",
					},
					{
						"name": "b",
						"wbs": "2",
						"amount": 1,
						"start_date": "2026-10-20",
						"end_date": "2026-10-21",
						"predecessors": [{"task": "a", "type": link_type}],
					},
				]
			)
			with self.subTest(link_type=link_type):
				self.assert_dates(rows["b"], start, end)
				self.assertEqual(rows["b"].duration, 2)

	def test_dependency_asap(self):
		# งานที่เริ่มช้ากว่าที่ความสัมพันธ์ยอมถูกดึงเข้ามา และลด lag แล้วงานขยับเร็วขึ้นตาม
		def plan(lag):
			return self.make_linked(
				[
					{
						"name": "a",
						"wbs": "1",
						"amount": 1,
						"start_date": "2026-11-02",
						"end_date": "2026-11-04",
					},
					{
						"name": "b",
						"wbs": "2",
						"amount": 1,
						"start_date": "2026-11-20",
						"end_date": "2026-11-21",
						"predecessors": [{"task": "a", "type": "FS", "lag": lag}],
					},
				]
			)

		self.assert_dates(plan(3)["b"], "2026-11-08", "2026-11-09")
		self.assert_dates(plan(0)["b"], "2026-11-05", "2026-11-06")

	def test_dependency_lag_and_working_days(self):
		# 7-8 พ.ย. เป็นเสาร์-อาทิตย์: จบศุกร์ 6 + FS+1 = เริ่มอังคาร 10 และ FS-1 = ซ้อน 1 วัน เริ่มศุกร์ 6
		weekend = ["2026-11-07", "2026-11-08"]
		rows = self.make_linked(
			[
				{"name": "a", "wbs": "1", "amount": 1, "start_date": "2026-11-05", "end_date": "2026-11-06"},
				{
					"name": "b",
					"wbs": "2",
					"amount": 1,
					"start_date": "2026-11-02",
					"duration": 2,
					"predecessors": [{"task": "a", "type": "FS", "lag": 1}],
				},
				{
					"name": "c",
					"wbs": "3",
					"amount": 1,
					"start_date": "2026-11-02",
					"duration": 2,
					"predecessors": [{"task": "a", "type": "FS", "lag": -1}],
				},
			],
			holidays=weekend,
		)
		self.assert_dates(rows["b"], "2026-11-10", "2026-11-11")
		self.assert_dates(rows["c"], "2026-11-06", "2026-11-09")
		self.assertEqual(rows["c"].duration, 2)

	def test_dependency_chain_and_missing_dates(self):
		# งานที่ตามมาอยู่ก่อนงานก่อนหน้าในตาราง ยังคำนวณตามลำดับความสัมพันธ์ และงานที่ไม่มีวันที่ได้วันเริ่มให้
		rows = self.make_linked(
			[
				{"name": "c", "wbs": "1", "amount": 1, "predecessors": [{"task": "b"}]},
				{
					"name": "b",
					"wbs": "2",
					"amount": 1,
					"start_date": "2026-11-02",
					"duration": 3,
					"predecessors": [{"task": "a"}, {"task": "deleted-row"}],
				},
				{"name": "a", "wbs": "3", "amount": 1, "start_date": "2026-11-02", "end_date": "2026-11-02"},
			]
		)
		self.assert_dates(rows["b"], "2026-11-03", "2026-11-05")
		self.assert_dates(rows["c"], "2026-11-06", "2026-11-06")
		# ลิงก์ที่ชี้ไปงานที่ไม่มีแล้วถูกตัดทิ้ง
		self.assertEqual(frappe.parse_json(rows["b"].predecessors), [{"task": "a", "type": "FS", "lag": 0}])

	def test_dependency_errors(self):
		cycle = [
			{
				"name": "a",
				"wbs": "1",
				"amount": 1,
				"start_date": "2026-11-02",
				"predecessors": [{"task": "b"}],
			},
			{
				"name": "b",
				"wbs": "2",
				"amount": 1,
				"start_date": "2026-11-02",
				"predecessors": [{"task": "a"}],
			},
		]
		with self.assertRaises(frappe.ValidationError):
			self.make_linked(cycle)

		on_group = [
			{"name": "g", "wbs": "1"},
			{"name": "a", "wbs": "1.1", "amount": 1, "start_date": "2026-11-02"},
			{
				"name": "b",
				"wbs": "2",
				"amount": 1,
				"start_date": "2026-11-02",
				"predecessors": [{"task": "g"}],
			},
		]
		with self.assertRaises(frappe.ValidationError):
			self.make_linked(on_group)

	def test_excel_round_trip(self):
		from frappe.utils.xlsxutils import read_xlsx_file_from_attached_file

		from consbiz.consbiz.doctype.construction_management.schedule_excel import (
			build_workbook,
			parse_schedule_rows,
		)

		self.make_linked(
			[
				{"name": "g", "wbs": "1", "item_name": "หมวด"},
				{
					"name": "a",
					"wbs": "1.1",
					"amount": 1,
					"start_date": "2026-11-02",
					"end_date": "2026-11-04",
				},
				{
					"name": "b",
					"wbs": "1.10",
					"amount": 1,
					"start_date": "2026-11-02",
					"duration": 2,
					"predecessors": [{"task": "a", "type": "SS", "lag": 2}],
				},
			]
		)
		content = build_workbook(self.linked_doc)
		parsed = parse_schedule_rows(read_xlsx_file_from_attached_file(fcontent=content))
		by_wbs = {row["wbs"]: row for row in parsed["rows"]}

		self.assertEqual(list(by_wbs), ["1", "1.1", "1.10"])
		self.assertEqual(by_wbs["1.10"]["predecessors"], "1.1SS+2")
		self.assertEqual(by_wbs["1.10"]["start_date"], "2026-11-04")
		self.assertEqual(by_wbs["1.10"]["end_date"], "2026-11-05")
		self.assertEqual(by_wbs["1.10"]["duration"], 2)
		self.assertEqual(by_wbs["1.1"]["predecessors"], "")
		self.assertFalse(parsed["warnings"])

	def test_excel_formulas(self):
		from io import BytesIO

		from openpyxl import load_workbook

		from consbiz.consbiz.doctype.construction_management.schedule_excel import (
			HEADER_ROW,
			build_workbook,
			parse_schedule_rows,
		)

		# 7-8 พ.ย. เป็นวันหยุด: งาน 1.2 จบวันอาทิตย์ 8 (สูตรจะได้ศุกร์ 6) จึงต้องเขียนเป็นค่า
		self.make_linked(
			[
				{"name": "g", "wbs": "1", "item_name": "หมวด"},
				{"name": "a", "wbs": "1.1", "amount": 1, "start_date": "2026-11-05", "duration": 3},
				{
					"name": "b",
					"wbs": "1.2",
					"amount": 1,
					"start_date": "2026-11-05",
					"end_date": "2026-11-08",
				},
			],
			holidays=["2026-11-07", "2026-11-08"],
		)
		content = build_workbook(self.linked_doc)
		wb = load_workbook(BytesIO(content))
		ws = wb.worksheets[0]
		group, task_a, task_b = (ws[HEADER_ROW + 2 + i] for i in range(3))
		end_col = 7  # คอลัมน์ H

		self.assertEqual(group[6].value, '=IF(COUNT(G7:G8)=0,"",MIN(G7:G8))')
		self.assertIn("WORKDAY.INTL(G7-1,MAX(I7,1),", task_a[end_col].value)
		self.assertNotIsInstance(task_b[end_col].value, str)
		self.assertIn("Holidays", wb.defined_names)
		self.assertEqual(len(wb.worksheets), 2)

		parsed = parse_schedule_rows(
			[
				[cell.value for cell in row]
				for row in load_workbook(BytesIO(content), data_only=True).active.iter_rows()
			]
		)
		by_wbs = {row["wbs"]: row for row in parsed["rows"]}
		self.assertEqual(by_wbs["1.1"]["end_date"], "2026-11-09")
		self.assertEqual(by_wbs["1.2"]["end_date"], "2026-11-08")
		self.assertEqual(by_wbs["1"]["start_date"], "2026-11-05")

	def test_excel_dates_mode(self):
		from io import BytesIO

		from openpyxl import load_workbook

		from consbiz.consbiz.doctype.construction_management.schedule_excel import (
			HEADER_ROW,
			build_workbook,
			parse_schedule_rows,
		)

		# แบบกำหนดวันเริ่ม-จบ: วันจบเป็นค่าที่กรอก จำนวนวันเป็นสูตรนับวันทำงาน (7-8 พ.ย. เป็นวันหยุด)
		self.make_linked(
			[{"name": "a", "wbs": "1", "amount": 1, "start_date": "2026-11-05", "end_date": "2026-11-10"}],
			holidays=["2026-11-07", "2026-11-08"],
		)
		content = build_workbook(self.linked_doc, mode="dates")
		row = load_workbook(BytesIO(content)).worksheets[0][HEADER_ROW + 2]
		self.assertNotIsInstance(row[7].value, str)
		self.assertIn("NETWORKDAYS.INTL(G6,H6,", row[8].value)

		parsed = parse_schedule_rows(
			[
				[cell.value for cell in r]
				for r in load_workbook(BytesIO(content), data_only=True).active.iter_rows()
			]
		)
		self.assertEqual(
			{k: parsed["rows"][0][k] for k in ("start_date", "end_date", "duration")},
			{"start_date": "2026-11-05", "end_date": "2026-11-10", "duration": 4},
		)

	def test_excel_parse_values(self):
		from consbiz.consbiz.doctype.construction_management.schedule_excel import (
			parse_date,
			parse_schedule_rows,
		)

		self.assertEqual(parse_date("15/01/2570"), "2027-01-15")
		self.assertEqual(parse_date("29/02/2567"), "2024-02-29")
		self.assertEqual(parse_date("2026-11-02"), "2026-11-02")
		self.assertIsNone(parse_date(""))
		with self.assertRaises(ValueError):
			parse_date("31/02/2026")

		parsed = parse_schedule_rows(
			[
				["แผนงาน"],
				["WBS", "Start", "Duration"],
				[2.0, "01/12/2026", "5 days"],
				[None, None, None],
				["3", "bad", ""],
			]
		)
		self.assertEqual(
			parsed["rows"],
			[
				{"row_no": 3, "wbs": "2", "start_date": "2026-12-01", "duration": 5},
				{"row_no": 5, "wbs": "3"},
			],
		)
		self.assertEqual(len(parsed["warnings"]), 1)

	def test_working_days(self):
		# 2026-11-07 และ 2026-11-08 เป็นเสาร์-อาทิตย์
		holidays = {getdate("2026-11-07"), getdate("2026-11-08")}
		self.assertEqual(count_working_days("2026-11-06", "2026-11-09", holidays), 2)
		self.assertEqual(count_working_days("2026-11-06", "2026-11-09", set()), 4)
		self.assertEqual(add_working_days("2026-11-06", 2, holidays), getdate("2026-11-09"))
		self.assertEqual(add_working_days("2026-11-07", 1, holidays), getdate("2026-11-09"))


class IntegrationTestConstructionProgress(IntegrationTestCase):
	def setUp(self):
		from consbiz.consbiz.doctype.construction_management.construction_management import create_from_boq

		boq = frappe.new_doc("BOQ")
		boq.naming_series = frappe.get_meta("BOQ").get_field("naming_series").options.split("\n")[0]
		boq.date = "2026-11-02"
		for row in (
			{"wbs": "1", "item_name": "หมวด", "quantity": 0},
			{"wbs": "1.1", "item_name": "งานมีปริมาณ", "quantity": 10, "total_amount": 3000},
			{"wbs": "1.2", "item_name": "งานเหมา", "quantity": 0, "total_amount": 1000},
		):
			boq.append("boq_details", row)
		boq.insert()
		self.doc = frappe.get_doc("Construction Management", create_from_boq(boq.name))
		self.group, self.qty_task, self.lump_task = self.doc.tasks

	def save(self, progress, tasks=None):
		from consbiz.consbiz.doctype.construction_management.construction_management import save_tasks

		self.doc.reload()
		data = save_tasks(
			self.doc.name, frappe.as_json(tasks or []), str(self.doc.modified), frappe.as_json(progress)
		)
		self.doc.reload()
		return data

	def test_progress_totals(self):
		self.save(
			[
				{"task": self.qty_task.name, "date": "2026-11-03", "qty": 4},
				{"task": self.qty_task.name, "date": "2026-11-04", "qty": 1},
				{"task": self.lump_task.name, "date": "2026-11-04", "percent": 20},
			]
		)
		qty_task, lump_task = self.doc.tasks[1], self.doc.tasks[2]
		self.assertEqual(qty_task.actual_qty, 5)
		self.assertAlmostEqual(qty_task.progress, 50)
		self.assertAlmostEqual(lump_task.progress, 20)
		# 75% x 50% + 25% x 20%
		self.assertAlmostEqual(self.doc.actual_progress, 42.5)
		self.assertAlmostEqual(self.doc.tasks[0].progress, 42.5)

		# ค่า 0 ลบรายการของวันนั้น
		self.save([{"task": self.qty_task.name, "date": "2026-11-04", "qty": 0}])
		self.assertEqual(self.doc.tasks[1].actual_qty, 4)

	def test_progress_cannot_exceed_quantity(self):
		with self.assertRaises(frappe.ValidationError):
			self.save([{"task": self.qty_task.name, "date": "2026-11-03", "qty": 11}])

	def save_baseline(self, title, activate=1):
		from consbiz.consbiz.doctype.construction_management.construction_management import save_baseline

		self.doc.reload()
		data = save_baseline(self.doc.name, str(self.doc.modified), title, activate=activate)
		self.doc.reload()
		return data["baselines"][-1]["name"]

	def test_baseline(self):
		from consbiz.consbiz.doctype.construction_management.construction_management import (
			activate_baseline,
			delete_baseline,
			target_progress,
		)

		first = self.save_baseline("แผนตามสัญญา")
		self.assertEqual(self.doc.active_baseline, first)
		self.assertEqual(self.doc.tasks[1].baseline_start_date, self.doc.tasks[1].start_date)
		self.assertEqual(self.doc.tasks[0].baseline_end_date, self.doc.tasks[0].end_date)
		self.assertTrue(self.doc.baseline_date)

		# เลื่อนงาน 1.1 ไป 10 พ.ย. แล้วบันทึกชุดที่ 2 โดยยังไม่ใช้เทียบ: ช่อง baseline_* ยังเป็นชุดแรก
		self.save([], tasks=[{"name": self.qty_task.name, "start_date": "2026-11-10", "end_date": "2026-11-10"}])
		second = self.save_baseline("ขยายเวลา EOT-1", activate=0)
		self.assertEqual(self.doc.active_baseline, first)
		self.assertEqual(str(self.doc.tasks[1].baseline_start_date), "2026-11-02")
		# 5 พ.ย. ชุดแรกครบ 100% ชุดที่ 2 มีแค่งาน 1.2 (น้ำหนัก 25%)
		self.assertAlmostEqual(target_progress(self.doc, "2026-11-05")[0], 100)

		activate_baseline(self.doc.name, str(self.doc.modified), second)
		self.doc.reload()
		self.assertEqual(self.doc.active_baseline, second)
		self.assertEqual(str(self.doc.tasks[1].baseline_start_date), "2026-11-10")
		self.assertAlmostEqual(target_progress(self.doc, "2026-11-05")[0], 25)

		# ชุดที่ใช้เทียบอยู่ลบไม่ได้ ชุดอื่นลบได้
		with self.assertRaises(frappe.ValidationError):
			delete_baseline(self.doc.name, second)
		self.assertEqual([b.name for b in delete_baseline(self.doc.name, first)], [second])

	def make_report(self, items, date="2026-11-05"):
		report = frappe.new_doc("Daily Report")
		report.construction_management = self.doc.name
		report.date = date
		for item in items:
			report.append("items", item)
		report.insert()
		return report

	def test_daily_report(self):
		report = self.make_report([{"wbs": "1.1", "qty": 3}, {"wbs": "1.2", "percent": 50}])
		self.assertAlmostEqual(report.items[0].percent, 30)
		self.assertEqual(report.items[0].remaining_qty, 7)
		report.submit()
		self.doc.reload()
		self.assertEqual(self.doc.tasks[1].actual_qty, 3)
		self.assertAlmostEqual(self.doc.tasks[2].progress, 50)

		# รวมกับฉบับก่อนแล้วเกินปริมาณ BOQ
		with self.assertRaises(frappe.ValidationError):
			self.make_report([{"wbs": "1.1", "qty": 8}], date="2026-11-06")
		# เลือกหัวข้อกลุ่มไม่ได้
		with self.assertRaises(frappe.ValidationError):
			self.make_report([{"wbs": "1", "percent": 10}], date="2026-11-06")

		# รายการจาก Daily Report ลบตรงไม่ได้ ต้อง Cancel
		entry = frappe.get_last_doc("Construction Progress Entry", {"reference_name": report.name})
		with self.assertRaises(frappe.ValidationError):
			entry.delete()

		report.cancel()
		self.doc.reload()
		self.assertEqual(self.doc.tasks[1].actual_qty, 0)
		self.assertFalse(frappe.db.exists("Construction Progress Entry", {"reference_name": report.name}))

	def test_daily_report_progress(self):
		# ทุกงานอยู่วันที่ 2 พ.ย. แผนสะสม ณ 5 พ.ย. = 100% ยังไม่ตั้ง Baseline จึงเทียบแผนปัจจุบัน
		# 1.1 ทำ 3/10 = 30% x น้ำหนัก 75% = 22.5%
		report = self.make_report([{"wbs": "1.1", "qty": 3}])
		self.assertAlmostEqual(report.today_progress, 22.5)
		self.assertAlmostEqual(report.cumulative_progress, 22.5)
		self.assertAlmostEqual(report.baseline_progress, 100)
		self.assertEqual(report.plan_basis, "Current Plan")
		self.assertAlmostEqual(report.progress_variance, -77.5)

		self.save_baseline("แผนตามสัญญา")
		report.save()
		self.assertEqual(report.plan_basis, "Baseline")
		report.submit()
		html = frappe.get_print("Daily Report", report.name, "Daily Report", doc=report)
		self.assertIn("ช้ากว่าแผน 77.500%", html)
		self.assertIn("แผนสะสม (Baseline)", html)

		# รายการหลังวันที่รายงานไม่นับในยอดสะสม 1.2 ทำ 50% x 25% = 12.5%
		self.save([{"task": self.lump_task.name, "date": "2026-11-10", "percent": 20}])
		later = self.make_report([{"wbs": "1.2", "percent": 50}], date="2026-11-06")
		self.assertAlmostEqual(later.today_progress, 12.5)
		self.assertAlmostEqual(later.cumulative_progress, 35)
		# Submit แล้วรายการของรายงานนี้ไม่ถูกนับซ้ำ
		later.submit()
		later.set_progress()
		self.assertAlmostEqual(later.cumulative_progress, 35)

	def test_plan_distribution(self):
		self.save(
			[],
			tasks=[
				{
					"name": self.qty_task.name,
					"start_date": "2026-11-01",
					"end_date": "2026-11-30",
					"plan_distribution": {"unit": "mweek", "weights": [10, 20, 30, 40]},
				}
			],
		)
		self.assertEqual(frappe.parse_json(self.doc.tasks[1].plan_distribution)["weights"], [10, 20, 30, 40])

		with self.assertRaises(frappe.ValidationError):
			self.save(
				[],
				tasks=[
					{
						"name": self.qty_task.name,
						"start_date": "2026-11-01",
						"end_date": "2026-11-30",
						"plan_distribution": {"unit": "mweek", "weights": [10, 20, 30, 30]},
					}
				],
			)

		# Baseline เก็บ % ที่กำหนดเองไปด้วย
		from consbiz.consbiz.doctype.construction_management.construction_management import set_baseline

		set_baseline(self.doc.name, str(self.doc.modified))
		self.doc.reload()
		self.assertEqual(self.doc.tasks[1].baseline_distribution, self.doc.tasks[1].plan_distribution)
